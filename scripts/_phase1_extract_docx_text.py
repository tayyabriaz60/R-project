"""Dump full paragraph/table text from each original docx for verification."""
from pathlib import Path
from docx import Document

root = Path(r"C:\Users\tayyab\Downloads\R-Project")
files = [
    root / "docs/spec/SAP Categorical Colormap Optimization.docx",
    root / "docs/spec/Final_Analysis_Implementation_Brief_Tayyab_FINAL.docx",
    root / "docs/spec/Data_Dictionary_Column_Structure_Studies_1_3_FINAL.docx",
]

out_dir = root / "docs" / "_docx_plaintext"
out_dir.mkdir(exist_ok=True)

for path in files:
    doc = Document(str(path))
    lines = [f"FILE: {path.name}", f"PARAGRAPHS: {len(doc.paragraphs)}", f"TABLES: {len(doc.tables)}", ""]
    lines.append("===== PARAGRAPHS =====")
    for i, p in enumerate(doc.paragraphs, 1):
        style = p.style.name if p.style is not None else ""
        lines.append(f"[{i}|{style}] {p.text}")
    lines.append("")
    lines.append("===== TABLES =====")
    for ti, t in enumerate(doc.tables, 1):
        lines.append(f"--- TABLE {ti} ({len(t.rows)}x{len(t.columns)}) ---")
        for ri, row in enumerate(t.rows):
            cells = [c.text.replace("\n", " | ") for c in row.cells]
            lines.append(f"R{ri}: " + " || ".join(cells))
        lines.append("")
    safe = path.stem.replace(" ", "_") + "_PLAINTEXT.txt"
    (out_dir / safe).write_text("\n".join(lines), encoding="utf-8")
    print(f"WROTE {safe} lines={len(lines)}")
