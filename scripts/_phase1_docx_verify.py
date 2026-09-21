"""Read-only Phase 1 verification: compare original .docx vs converted .md."""
from __future__ import annotations

import json
import re
import zipfile
from collections import Counter
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.oxml.ns import qn

NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "wps": "http://schemas.microsoft.com/office/word/2010/wordprocessingShape",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "w14": "http://schemas.microsoft.com/office/word/2010/wordml",
}


def count_md(path: Path) -> dict:
    text = path.read_text(encoding="utf-8")
    lines = text.splitlines()
    headings = [ln for ln in lines if re.match(r"^#{1,6}\s", ln)]
    # GFM tables: header row then separator
    table_seps = [i for i, ln in enumerate(lines) if re.match(r"^\|?\s*:?-{3,}", ln) and "|" in ln]
    list_items = [ln for ln in lines if re.match(r"^\s*([-*+]|\d+\.)\s+", ln)]
    images = re.findall(r"!\[[^\]]*\]\([^)]+\)", text)
    footnotes = re.findall(r"\[\^[^\]]+\]", text)
    html_comments = re.findall(r"<!--.*?-->", text, flags=re.S)
    return {
        "chars": len(text),
        "lines": len(lines),
        "headings": len(headings),
        "heading_texts": headings,
        "tables_est": len(table_seps),
        "list_items": len(list_items),
        "images": len(images),
        "image_refs": images,
        "footnote_markers": len(footnotes),
        "html_comments": len(html_comments),
        "pipe_rows": sum(1 for ln in lines if ln.strip().startswith("|")),
    }


def paragraph_style(p) -> str:
    try:
        return p.style.name if p.style is not None else ""
    except Exception:
        return ""


def inspect_docx(path: Path) -> dict:
    doc = Document(str(path))
    heading_texts = []
    list_paras = 0
    body_paras = 0
    empty_paras = 0
    for p in doc.paragraphs:
        body_paras += 1
        style = paragraph_style(p)
        text = p.text.strip()
        if not text:
            empty_paras += 1
        if style.lower().startswith("heading") or re.match(r"^Heading\s*\d", style, re.I):
            heading_texts.append(f"{style}: {text}")
        pPr = p._p.find(qn("w:pPr"))
        if pPr is not None and pPr.find(qn("w:numPr")) is not None:
            list_paras += 1

    tables = []
    for i, t in enumerate(doc.tables, 1):
        nrows = len(t.rows)
        ncols = len(t.columns) if t.rows else 0
        # cell texts first row
        header = []
        if t.rows:
            header = [c.text.strip().replace("\n", " ") for c in t.rows[0].cells]
        tables.append({"index": i, "rows": nrows, "cols": ncols, "header": header})

    # Raw XML for comments, revisions, footnotes, images, headers
    raw = {
        "comments": 0,
        "comment_texts": [],
        "footnotes": 0,
        "footnote_texts": [],
        "endnotes": 0,
        "ins": 0,
        "del": 0,
        "images_rels": 0,
        "drawing_blips": 0,
        "headers": [],
        "footers": [],
        "hyperlinks": 0,
        "sectPr": 0,
    }
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        raw["zip_entries"] = names
        if "word/comments.xml" in names:
            root = ET.fromstring(z.read("word/comments.xml"))
            comments = root.findall(".//w:comment", NS)
            raw["comments"] = len(comments)
            for c in comments:
                texts = [t.text or "" for t in c.findall(".//w:t", NS)]
                raw["comment_texts"].append("".join(texts))
        if "word/footnotes.xml" in names:
            root = ET.fromstring(z.read("word/footnotes.xml"))
            notes = [
                n
                for n in root.findall("w:footnote", NS)
                if n.get(qn("w:type")) not in ("separator", "continuationSeparator")
            ]
            raw["footnotes"] = len(notes)
            for n in notes:
                texts = [t.text or "" for t in n.findall(".//w:t", NS)]
                raw["footnote_texts"].append("".join(texts))
        if "word/endnotes.xml" in names:
            root = ET.fromstring(z.read("word/endnotes.xml"))
            notes = [
                n
                for n in root.findall("w:endnote", NS)
                if n.get(qn("w:type")) not in ("separator", "continuationSeparator")
            ]
            raw["endnotes"] = len(notes)
        document_xml = z.read("word/document.xml")
        droot = ET.fromstring(document_xml)
        raw["ins"] = len(droot.findall(".//w:ins", NS))
        raw["del"] = len(droot.findall(".//w:del", NS))
        raw["drawing_blips"] = len(droot.findall(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip"))
        raw["hyperlinks"] = len(droot.findall(".//w:hyperlink", NS))
        raw["sectPr"] = len(droot.findall(".//w:sectPr", NS))
        # relationships for images
        if "word/_rels/document.xml.rels" in names:
            rels = ET.fromstring(z.read("word/_rels/document.xml.rels"))
            for rel in rels:
                tgt = rel.get("Target") or ""
                typ = rel.get("Type") or ""
                if "image" in typ.lower() or tgt.startswith("media/"):
                    raw["images_rels"] += 1
        media = [n for n in names if n.startswith("word/media/")]
        raw["media_files"] = media
        raw["media_count"] = len(media)
        for n in names:
            if n.startswith("word/header"):
                hroot = ET.fromstring(z.read(n))
                texts = [t.text or "" for t in hroot.findall(".//w:t", NS)]
                raw["headers"].append({"file": n, "text": "".join(texts)})
            if n.startswith("word/footer"):
                froot = ET.fromstring(z.read(n))
                texts = [t.text or "" for t in froot.findall(".//w:t", NS)]
                raw["footers"].append({"file": n, "text": "".join(texts)})

    styles = Counter(paragraph_style(p) or "(none)" for p in doc.paragraphs)
    core = {}
    try:
        cp = doc.core_properties
        core = {
            "title": cp.title,
            "author": cp.author,
            "last_modified_by": cp.last_modified_by,
            "created": str(cp.created),
            "modified": str(cp.modified),
            "revision": cp.revision,
        }
    except Exception as e:
        core = {"error": str(e)}

    return {
        "file": path.name,
        "bytes": path.stat().st_size,
        "core": core,
        "paragraphs": body_paras,
        "empty_paragraphs": empty_paras,
        "headings": len(heading_texts),
        "heading_texts": heading_texts,
        "list_paragraphs": list_paras,
        "tables": len(tables),
        "table_details": tables,
        "styles": dict(styles),
        "xml": {
            k: v
            for k, v in raw.items()
            if k != "zip_entries"
        },
        "zip_entry_count": len(raw["zip_entries"]),
    }


def main() -> None:
    root = Path(r"C:\Users\tayyab\Downloads\R-Project")
    pairs = [
        (
            root / "docs/spec/Data_Dictionary_Column_Structure_Studies_1_3_FINAL.docx",
            root / "docs/Data_Dictionary_Column_Structure_Studies_1_3_FINAL.md",
        ),
        (
            root / "docs/spec/Final_Analysis_Implementation_Brief_Tayyab_FINAL.docx",
            root / "docs/Final_Analysis_Implementation_Brief_Tayyab_FINAL.md",
        ),
        (
            root / "docs/spec/SAP Categorical Colormap Optimization.docx",
            root / "docs/SAP_Categorical_Colormap_Optimization.md",
        ),
    ]
    reports = []
    for docx_path, md_path in pairs:
        d = inspect_docx(docx_path)
        m = count_md(md_path)
        issues = []
        if d["tables"] != m["tables_est"]:
            issues.append(f"TABLE COUNT mismatch: docx={d['tables']} md_sep_rows={m['tables_est']} md_pipe_rows={m['pipe_rows']}")
        if d["headings"] != m["headings"]:
            issues.append(f"HEADING COUNT mismatch: docx={d['headings']} md={m['headings']}")
        if d["xml"]["footnotes"] and m["footnote_markers"] == 0:
            issues.append(f"FOOTNOTES in docx ({d['xml']['footnotes']}) but no MD footnote markers")
        if d["xml"]["comments"]:
            issues.append(f"COMMENTS in docx ({d['xml']['comments']}) — pandoc GFM typically DROPS comments")
        if d["xml"]["ins"] or d["xml"]["del"]:
            issues.append(f"TRACKED CHANGES in docx ins={d['xml']['ins']} del={d['xml']['del']}")
        if d["xml"]["media_count"] and m["images"] == 0:
            issues.append(f"IMAGES in docx media={d['xml']['media_count']} but no MD image refs")
        if d["list_paragraphs"] and m["list_items"] == 0:
            issues.append(f"LISTS in docx ({d['list_paragraphs']}) but no MD list items")
        reports.append({"docx": d, "md": m, "issues": issues})

    out = root / "docs" / "_conversion_verification.json"
    out.write_text(json.dumps(reports, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(reports, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
